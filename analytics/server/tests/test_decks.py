import base64

import pandas as pd
import pytest
from pptx import Presentation

from vogelring_analytics.datasets import DatasetStore
from vogelring_analytics.decks import DeckBuildError, DeckSpec, DeckStore, Slide, build_deck, deck_from_report, deck_urls
from vogelring_analytics.visuals import VisualStore

# 1×1 red PNG
_PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFBQIAX8jx0gAAAABJRU5ErkJggg==")


@pytest.fixture
def stores(tmp_path):
    ds = DatasetStore(tmp_path / "ds")
    vs = VisualStore(tmp_path / "viz")
    dk = DeckStore(tmp_path / "decks")
    meta = ds.save(pd.DataFrame({"ring": ["A1", "B2"], "ort": ["Ostpark", "Hafen"], "sichtungen": [12, 7]}), name="Top-Orte", sql="SELECT 1", description="Alle Sichtungen 2025.")
    vm = vs.create("chart", {"kind": "bar"}, {"kind": "bar", "width": 1200, "height": 700}, title="Sichtungen je Ort", dataset_ids=[meta.id], width=1200, height=700, summary={})
    vs.write_bytes(vm.id, "image.png", _PNG)
    missing = vs.create("map", {}, {"width": 1400, "height": 900}, title="Karte ohne Bild", dataset_ids=[], width=1400, height=900, summary={})
    return ds, vs, dk, meta, vm, missing


def test_build_deck_all_kinds(stores):
    ds, vs, dk, meta, vm, missing = stores
    spec = DeckSpec(
        title="Ostpark 2025", subtitle="Mauserzeit", question="Wer war da?",
        slides=[
            Slide(kind="key_findings", bullets=["**109** Vögel", "Graugänse dominieren", "Juni ist der Höhepunkt", "vier ohne Koordinaten"]),
            Slide(kind="section", title="Ergebnisse", subtitle="Was wurde gezählt"),
            Slide(kind="bullets", title="Punkte", bullets=[f"Punkt {i}" for i in range(9)], notes="Sprechernotiz"),
            Slide(kind="figure", visual_ids=[vm.id], caption="Abbildung 1"),
            Slide(kind="two_figures", title="Zwei Bilder", visual_ids=[vm.id, missing.id]),
            Slide(kind="table", dataset_id=meta.id, caption="Tabelle 1"),
            Slide(kind="text", title="Einordnung", text="Erster Absatz.\n\nZweiter **wichtiger** Absatz."),
        ],
        closing_text="Danke",
    )
    dmeta = build_deck(spec, visuals=vs, datasets=ds, store=dk)
    path = dk.file(dmeta.id, "deck.pptx")
    assert path and path.exists() and path.stat().st_size > 20_000
    prs = Presentation(str(path))
    # title + 7 content + methods + closing
    assert len(prs.slides) == 10 and dmeta.slides == 10
    assert dmeta.visual_ids == [vm.id, missing.id] and meta.id in dmeta.dataset_ids
    assert dmeta.files == deck_urls(dmeta.id) and dmeta.files["pptx_url"].endswith(f"/decks/{dmeta.id}.pptx")
    assert prs.slides[3].has_notes_slide and "Sprechernotiz" in prs.slides[3].notes_slide.notes_text_frame.text
    assert dk.exists(dmeta.id) and dk.list()[0].id == dmeta.id and dk.meta(dmeta.id).title == "Ostpark 2025"


def test_missing_image_placeholder_and_errors(stores):
    ds, vs, dk, meta, vm, missing = stores
    spec = DeckSpec(title="T", slides=[Slide(kind="figure", visual_ids=[missing.id])], include_methods=False)
    dmeta = build_deck(spec, visuals=vs, datasets=ds, store=dk)
    prs = Presentation(str(dk.file(dmeta.id, "deck.pptx")))
    texts = " ".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
    assert "Abbildung nicht verfügbar" in texts and len(prs.slides) == 3
    with pytest.raises(DeckBuildError):
        build_deck(DeckSpec(title="T", slides=[Slide(kind="figure")]), visuals=vs, datasets=ds, store=dk)
    with pytest.raises(DeckBuildError):
        build_deck(DeckSpec(title="T", slides=[Slide(kind="table", dataset_id="ds_nope")]), visuals=vs, datasets=ds, store=dk)


def test_deck_from_report(stores):
    ds, vs, dk, meta, vm, missing = stores
    report = {
        "title": "Bericht", "subtitle": "Unter", "question": "Frage?", "author": "Ingo", "date": "01.02.2026",
        "key_findings": ["A", "B"],
        "sections": [{"heading": "Ergebnisse", "text": "Erster Satz. Zweiter Satz.\n\n- eins\n- zwei", "blocks": [{"type": "figure", "id": vm.id, "no": 1, "caption": "Bild"}, {"type": "table", "no": 1, "caption": "Tab", "dataset_id": meta.id, "model": {}}]}],
        "recommendations": ["Weiter beobachten."],
    }
    spec = deck_from_report(report)
    kinds = [s.kind for s in spec.slides]
    assert kinds == ["key_findings", "section", "bullets", "figure", "table", "bullets"]
    assert spec.slides[1].subtitle == "Erster Satz." and spec.slides[2].bullets == ["eins", "zwei"]
    dmeta = build_deck(spec, visuals=vs, datasets=ds, store=dk)
    assert len(Presentation(str(dk.file(dmeta.id, "deck.pptx"))).slides) == 9
