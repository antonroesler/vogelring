import pandas as pd
from openpyxl import load_workbook

from vogelring_analytics.exporters import safe_filename, write_csv, write_xlsx


def test_xlsx(tmp_path):
    df = pd.DataFrame({"ring": ["A"], "date": pd.to_datetime(["2025-01-02"]), "n": [3]})
    p = write_xlsx(df, tmp_path / "x.xlsx", title="Test")
    ws = load_workbook(p).active
    assert ws["A3"].value == "ring" and ws["B4"].number_format == "DD.MM.YYYY" and ws["C4"].value == 3


def test_csv_and_names(tmp_path):
    df = pd.DataFrame({"a": [1]})
    assert write_csv(df, tmp_path / "a.csv").read_text(encoding="utf-8-sig").startswith("a\n")
    assert safe_filename("Wiederfunde: Ostpark/2025") == "Wiederfunde_Ostpark_2025"
    assert safe_filename("") == "export"
