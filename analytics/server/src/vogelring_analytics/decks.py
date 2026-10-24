"""Presentation composer: DeckSpec (German slides from existing visuals/datasets) → PPTX in the house style."""

from __future__ import annotations

import json
import re
import threading
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt
from pydantic import BaseModel, ConfigDict, Field

from . import palette
from .datasets import DatasetStore, new_id
from .events import bus
from .settings import settings
from .specs import TableSpec
from .tables import build_table_model
from .visuals import VisualStore

SlideKind = Literal["title", "section", "bullets", "figure", "two_figures", "table", "key_findings", "text", "methods"]


class DeckBuildError(ValueError):
    pass


class _Spec(BaseModel):
    model_config = ConfigDict(extra="forbid")


class Slide(_Spec):
    kind: SlideKind
    title: str | None = Field(None, description="Folientitel (Deutsch).")
    subtitle: str | None = None
    bullets: list[str] = Field(default_factory=list, description="Stichpunkte; **fett** wird unterstützt. Max. 7 je Folie.")
    text: str | None = Field(None, description="Fließtext (kind=text), Absätze durch Leerzeile.")
    visual_ids: list[str] = Field(default_factory=list, description="figure: 1 visual_id, two_figures: 2.")
    dataset_id: str | None = Field(None, description="table: Datensatz (max. 15 Zeilen × 8 Spalten).")
    caption: str | None = Field(None, description="Bildunterschrift.")
    notes: str | None = Field(None, description="Sprechernotizen in einfachem Deutsch.")


class DeckSpec(_Spec):
    title: str
    subtitle: str | None = None
    author: str = "Vogelring-Analyse"
    date: str | None = Field(None, description="dd.mm.yyyy; Standard: heute.")
    question: str | None = Field(None, description="Die Fragestellung in einem Satz (Kasten auf der Titelfolie).")
    slides: list[Slide] = Field(..., min_length=1)
    include_methods: bool = Field(True, description="Abschlussfolie „So wurde gezählt“ aus den Datensatz-Beschreibungen.")
    closing_text: str | None = Field(None, description="Text der Schlussfolie, z. B. Dank oder nächste Schritte.")


@dataclass
class DeckMeta:
    id: str
    title: str
    created_at: str
    slides: int
    visual_ids: list[str]
    dataset_ids: list[str]
    files: dict[str, str] = field(default_factory=dict)


def deck_urls(deck_id: str) -> dict[str, str]:
    return {"pptx_url": f"{settings.public_base_url}/decks/{deck_id}.pptx"}


class DeckStore:
    def __init__(self, root: Path | None = None) -> None:
        self.root = root or (settings.data_dir / "decks")
        self.root.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()

    def dir(self, deck_id: str) -> Path:
        if not deck_id or any(ch in deck_id for ch in "/\\.."):
            raise KeyError(f"Ungültige Präsentations-ID: {deck_id!r}")
        return self.root / deck_id

    def create(self, spec: dict[str, Any], meta_kwargs: dict[str, Any]) -> DeckMeta:
        deck_id = new_id("dk_")
        d = self.dir(deck_id)
        d.mkdir(parents=True, exist_ok=False)
        meta = DeckMeta(id=deck_id, created_at=datetime.now(timezone.utc).isoformat(timespec="seconds"), **meta_kwargs)
        with self._lock:
            (d / "spec.json").write_text(json.dumps(spec, ensure_ascii=False, indent=1, default=str))
            (d / "meta.json").write_text(json.dumps(asdict(meta), ensure_ascii=False, indent=1))
        return meta

    def save_meta(self, meta: DeckMeta) -> None:
        with self._lock:
            (self.dir(meta.id) / "meta.json").write_text(json.dumps(asdict(meta), ensure_ascii=False, indent=1))

    def meta(self, deck_id: str) -> DeckMeta:
        p = self.dir(deck_id) / "meta.json"
        if not p.exists():
            raise KeyError(f"Unbekannte Präsentations-ID: {deck_id}")
        return DeckMeta(**json.loads(p.read_text()))

    def exists(self, deck_id: str) -> bool:
        try:
            return (self.dir(deck_id) / "meta.json").exists()
        except KeyError:
            return False

    def file(self, deck_id: str, name: str) -> Path | None:
        p = self.dir(deck_id) / name
        return p if p.exists() else None

    def list(self, limit: int = 20) -> list[DeckMeta]:
        metas = []
        for p in self.root.glob("dk_*/meta.json"):
            try:
                metas.append(DeckMeta(**json.loads(p.read_text())))
            except Exception:
                continue
        metas.sort(key=lambda m: m.created_at, reverse=True)
        return metas[:limit]


# ---------------------------------------------------------------- styling helpers

PRIMARY = RGBColor.from_string(palette.THEME["primary"][1:])
SECONDARY = RGBColor.from_string(palette.THEME["secondary"][1:])
ACCENT = RGBColor.from_string(palette.THEME["accent"][1:])
BACKGROUND = RGBColor.from_string(palette.THEME["background"][1:])
TEXT = RGBColor(0x1F, 0x2D, 0x36)
MUTED = RGBColor.from_string(palette.THEME["muted"][1:])
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ZEBRA = RGBColor(0xF3, 0xF4, 0xF6)
FONT = "Calibri"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
HEADER_H = Inches(1.05)
FOOTER_H = Inches(0.4)
CONTENT_TOP = HEADER_H + Inches(0.1)
CONTENT_H = SLIDE_H - CONTENT_TOP - FOOTER_H - Inches(0.1)
CONTENT_W = SLIDE_W - 2 * MARGIN

_BOLD = re.compile(r"(\*\*[^*]+\*\*)")


def _runs(paragraph, text: str, *, size: int, color: RGBColor = TEXT, bold: bool = False) -> None:
    """Fill a paragraph with runs; **x** becomes bold."""
    parts = [p for p in _BOLD.split(text or "") if p]
    if not parts:
        parts = [""]
    for part in parts:
        run = paragraph.add_run()
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run.text = part[2:-2]
            run.font.bold = True
        else:
            run.text = part
            run.font.bold = bold
        run.font.size = Pt(size)
        run.font.name = FONT
        run.font.color.rgb = color


def _textbox(slide, left, top, width, height, text: str, *, size: int, color: RGBColor = TEXT, bold: bool = False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap: bool = True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paragraphs = (text or "").split("\n")
    for i, line in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _runs(p, line, size=size, color=color, bold=bold)
    return box


def _rect(slide, left, top, width, height, fill: RGBColor, line: RGBColor | None = None):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def _blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _header(slide, title: str | None, subtitle: str | None = None) -> None:
    _rect(slide, 0, 0, SLIDE_W, Inches(0.08), SECONDARY)
    _textbox(slide, MARGIN, Inches(0.22), CONTENT_W, Inches(0.55), title or "", size=26, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        _textbox(slide, MARGIN, Inches(0.74), CONTENT_W, Inches(0.3), subtitle, size=13, color=MUTED)


def _footer(slide, date: str, number: int) -> None:
    top = SLIDE_H - FOOTER_H
    _textbox(slide, MARGIN, top, Inches(5), FOOTER_H, f"Vogelring · {date}", size=10, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    _textbox(slide, SLIDE_W - MARGIN - Inches(1.2), top, Inches(1.2), FOOTER_H, str(number), size=10, color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def _notes(slide, text: str | None) -> None:
    if text:
        slide.notes_slide.notes_text_frame.text = text


def _png_size(data: bytes) -> tuple[int, int] | None:
    if len(data) >= 24 and data[:8] == b"\x89PNG\r\n\x1a\n":
        return int.from_bytes(data[16:20], "big"), int.from_bytes(data[20:24], "big")
    return None


def _fit(img_w: int, img_h: int, box_w: int, box_h: int) -> tuple[int, int]:
    scale = min(box_w / img_w, box_h / img_h)
    return int(img_w * scale), int(img_h * scale)


# ---------------------------------------------------------------- slide builders


class _Builder:
    def __init__(self, spec: DeckSpec, visuals: VisualStore, datasets: DatasetStore) -> None:
        self.spec = spec
        self.visuals = visuals
        self.datasets = datasets
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        self.date = spec.date or datetime.now().strftime("%d.%m.%Y")
        self.number = 0
        self.visual_ids: list[str] = []
        self.dataset_ids: list[str] = []

    # -- generic content slide scaffold
    def content_slide(self, title: str | None, subtitle: str | None = None):
        slide = _blank_slide(self.prs)
        self.number += 1
        _header(slide, title, subtitle)
        _footer(slide, self.date, self.number)
        return slide

    def title_slide(self) -> None:
        s = self.spec
        slide = _blank_slide(self.prs)
        self.number += 1
        _rect(slide, 0, 0, SLIDE_W, SLIDE_H, BACKGROUND)
        _rect(slide, 0, 0, Inches(0.35), SLIDE_H, PRIMARY)
        _textbox(slide, Inches(1.0), Inches(1.1), CONTENT_W, Inches(0.4), "VOGELRING · AUSWERTUNG", size=12, color=SECONDARY, bold=True)
        _textbox(slide, Inches(1.0), Inches(1.6), SLIDE_W - Inches(2.0), Inches(1.9), s.title, size=40, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        y = Inches(3.55)
        if s.subtitle:
            _textbox(slide, Inches(1.0), y, SLIDE_W - Inches(2.0), Inches(0.7), s.subtitle, size=20, color=MUTED)
            y += Inches(0.75)
        _textbox(slide, Inches(1.0), y, SLIDE_W - Inches(2.0), Inches(0.4), f"{s.author} · {self.date}", size=14, color=MUTED)
        if s.question:
            top = Inches(5.3)
            _rect(slide, Inches(1.0), top, SLIDE_W - Inches(2.0), Inches(1.2), WHITE, line=RGBColor(0xE3, 0xE8, 0xEC))
            _rect(slide, Inches(1.0), top, Inches(0.08), Inches(1.2), SECONDARY)
            _textbox(slide, Inches(1.25), top + Inches(0.1), SLIDE_W - Inches(2.5), Inches(1.0), f"**Fragestellung:** {s.question}", size=15, anchor=MSO_ANCHOR.MIDDLE)

    def section_slide(self, sl: Slide) -> None:
        slide = _blank_slide(self.prs)
        self.number += 1
        _rect(slide, 0, 0, SLIDE_W, SLIDE_H, PRIMARY)
        _rect(slide, MARGIN, Inches(3.0), Inches(1.2), Inches(0.08), ACCENT)
        _textbox(slide, MARGIN, Inches(3.2), CONTENT_W, Inches(1.4), sl.title or "", size=36, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        if sl.subtitle:
            _textbox(slide, MARGIN, Inches(4.6), CONTENT_W, Inches(0.8), sl.subtitle, size=18, color=RGBColor(0xD6, 0xE4, 0xEC))
        _textbox(slide, SLIDE_W - MARGIN - Inches(1.2), SLIDE_H - FOOTER_H, Inches(1.2), FOOTER_H, str(self.number), size=10, color=RGBColor(0xD6, 0xE4, 0xEC), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        _notes(slide, sl.notes)

    def bullets_slide(self, sl: Slide) -> None:
        slide = self.content_slide(sl.title, sl.subtitle)
        items = sl.bullets[:7]
        size = 22 if len(items) <= 4 else (19 if len(items) <= 6 else 17)
        box = slide.shapes.add_textbox(MARGIN, CONTENT_TOP + Inches(0.2), CONTENT_W, CONTENT_H - Inches(0.3))
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            _runs(p, "•  " + item, size=size)
        if len(sl.bullets) > 7:
            p = tf.add_paragraph()
            _runs(p, f"… und {len(sl.bullets) - 7} weitere Punkte", size=12, color=MUTED)
        _notes(slide, sl.notes)

    def text_slide(self, sl: Slide) -> None:
        slide = self.content_slide(sl.title, sl.subtitle)
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", sl.text or "") if p.strip()]
        box = slide.shapes.add_textbox(MARGIN, CONTENT_TOP + Inches(0.2), CONTENT_W, CONTENT_H - Inches(0.3))
        tf = box.text_frame
        tf.word_wrap = True
        total = sum(len(p) for p in paragraphs)
        size = 18 if total < 600 else (16 if total < 1000 else 14)
        for i, para in enumerate(paragraphs or [""]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            _runs(p, para.replace("\n", " "), size=size)
        _notes(slide, sl.notes)

    def _figure(self, slide, visual_id: str, left: int, top: int, box_w: int, box_h: int, caption: str | None, *, caption_size: int = 12) -> None:
        self.visual_ids.append(visual_id)
        png = self.visuals.file(visual_id, "image.png") if self.visuals.exists(visual_id) else None
        cap_h = Inches(0.42) if caption else Inches(0.26)
        img_box_h = box_h - cap_h - Inches(0.05)
        if png:
            data = png.read_bytes()
            size = _png_size(data)
            w, h = _fit(size[0], size[1], box_w, img_box_h) if size else (box_w, img_box_h)
            slide.shapes.add_picture(str(png), left + (box_w - w) // 2, top, width=Emu(w), height=Emu(h))
            bottom = top + h
        else:
            title = self.visuals.meta(visual_id).title if self.visuals.exists(visual_id) else visual_id
            _rect(slide, left, top, box_w, img_box_h, ZEBRA, line=RGBColor(0xE3, 0xE8, 0xEC))
            _textbox(slide, left, top, box_w, img_box_h, f"Abbildung nicht verfügbar\n{title}", size=14, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            bottom = top + img_box_h
        if caption:
            _textbox(slide, left, bottom + Inches(0.03), box_w, Inches(0.3), caption, size=caption_size, color=TEXT, align=PP_ALIGN.CENTER)
            bottom += Inches(0.3)
        _textbox(slide, left, bottom, box_w, Inches(0.22), f"Interaktiv: {settings.public_base_url}/viz/{visual_id}", size=8, color=MUTED, align=PP_ALIGN.CENTER)

    def figure_slide(self, sl: Slide) -> None:
        if not sl.visual_ids:
            raise DeckBuildError("Folie 'figure' braucht eine visual_id.")
        vid = sl.visual_ids[0]
        title = sl.title or (self.visuals.meta(vid).title if self.visuals.exists(vid) else None)
        slide = self.content_slide(title, sl.subtitle)
        self._figure(slide, vid, MARGIN, CONTENT_TOP, CONTENT_W, CONTENT_H, sl.caption)
        _notes(slide, sl.notes)

    def two_figures_slide(self, sl: Slide) -> None:
        if len(sl.visual_ids) < 2:
            raise DeckBuildError("Folie 'two_figures' braucht zwei visual_ids.")
        slide = self.content_slide(sl.title, sl.subtitle)
        gap = Inches(0.3)
        w = (CONTENT_W - gap) // 2
        for i, vid in enumerate(sl.visual_ids[:2]):
            cap = (self.visuals.meta(vid).title if self.visuals.exists(vid) else None) if sl.caption is None else (sl.caption if i == 0 else None)
            self._figure(slide, vid, MARGIN + i * (w + gap), CONTENT_TOP, w, CONTENT_H, cap, caption_size=11)
        _notes(slide, sl.notes)

    def table_slide(self, sl: Slide) -> None:
        if not sl.dataset_id:
            raise DeckBuildError("Folie 'table' braucht eine dataset_id.")
        self.dataset_ids.append(sl.dataset_id)
        df = self.datasets.load(sl.dataset_id)
        meta = self.datasets.meta(sl.dataset_id)
        model, _ = build_table_model(TableSpec(dataset_id=sl.dataset_id, limit=15, row_number=False), df)
        cols = model["columns"][:8]
        rows = [r[: len(cols)] for r in model["rows"]]
        slide = self.content_slide(sl.title or meta.name, sl.subtitle)
        n_rows = len(rows) + 1
        row_h = min(Inches(0.38), int((CONTENT_H - Inches(0.6)) / max(n_rows, 1)))
        shape = slide.shapes.add_table(n_rows, len(cols), MARGIN, CONTENT_TOP + Inches(0.05), CONTENT_W, row_h * n_rows)
        table = shape.table
        font_size = 13 if n_rows <= 8 else (11 if n_rows <= 12 else 10)
        for j, c in enumerate(cols):
            cell = table.cell(0, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PRIMARY
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if c["align"] == "right" else PP_ALIGN.LEFT
            _runs(p, c["label"], size=font_size, color=WHITE, bold=True)
        for i, row in enumerate(rows, start=1):
            for j, cell_model in enumerate(row):
                cell = table.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = ZEBRA if i % 2 == 0 else WHITE
                cell.text = ""
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT if cols[j]["align"] == "right" else PP_ALIGN.LEFT
                _runs(p, cell_model["text"], size=font_size, bold=bool(cell_model.get("hl")))
        for i in range(n_rows):
            table.rows[i].height = row_h
        bottom = CONTENT_TOP + Inches(0.05) + row_h * n_rows + Inches(0.08)
        note = sl.caption or ""
        if model["total_rows"] > model["shown_rows"]:
            note = (note + " · " if note else "") + f"{model['shown_rows']} von {model['total_rows']} Zeilen gezeigt"
        if len(model["columns"]) > 8:
            note = (note + " · " if note else "") + f"{len(model['columns']) - 8} Spalten ausgeblendet"
        if note:
            _textbox(slide, MARGIN, bottom, CONTENT_W, Inches(0.3), note, size=11, color=MUTED)
        _notes(slide, sl.notes)

    def key_findings_slide(self, sl: Slide) -> None:
        slide = self.content_slide(sl.title or "Das Wichtigste in Kürze", sl.subtitle)
        items = sl.bullets[:6]
        if not items:
            raise DeckBuildError("Folie 'key_findings' braucht bullets.")
        cols = 1 if len(items) <= 3 else 2
        rows = -(-len(items) // cols)
        gap = Inches(0.25)
        card_w = (CONTENT_W - gap * (cols - 1)) // cols
        card_h = min(Inches(1.5), int((CONTENT_H - gap * (rows - 1)) / rows))
        size = 18 if cols == 1 else 15
        for i, item in enumerate(items):
            r, c = divmod(i, cols)
            left = MARGIN + c * (card_w + gap)
            top = CONTENT_TOP + r * (card_h + gap)
            _rect(slide, left, top, card_w, card_h, WHITE, line=RGBColor(0xE3, 0xE8, 0xEC))
            _rect(slide, left, top, Inches(0.08), card_h, ACCENT)
            _textbox(slide, left + Inches(0.2), top, Inches(0.7), card_h, str(i + 1), size=30, color=SECONDARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            _textbox(slide, left + Inches(0.9), top + Inches(0.08), card_w - Inches(1.05), card_h - Inches(0.16), item, size=size, anchor=MSO_ANCHOR.MIDDLE)
        _notes(slide, sl.notes)

    def methods_slide(self, sl: Slide | None = None) -> None:
        slide = self.content_slide((sl.title if sl and sl.title else None) or "So wurde gezählt", sl.subtitle if sl else None)
        ids = list(dict.fromkeys(self.dataset_ids))
        box = slide.shapes.add_textbox(MARGIN, CONTENT_TOP + Inches(0.1), CONTENT_W, CONTENT_H - Inches(0.2))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.space_after = Pt(10)
        _runs(p, "Grundlage sind die Sichtungen und Beringungen in der Vogelring-Datenbank. Jede Sichtung ist eine Ablesung eines beringten Vogels an einem Ort an einem Tag. Nicht gesehen heißt nicht abwesend.", size=14, color=MUTED)
        size = 14 if len(ids) <= 4 else 12
        for ds_id in ids[:8]:
            try:
                meta = self.datasets.meta(ds_id)
            except KeyError:
                continue
            p = tf.add_paragraph()
            p.space_after = Pt(8)
            desc = f" – {meta.description}" if meta.description else ""
            _runs(p, f"**{meta.name}**{desc} ({meta.row_count} Zeilen)", size=size)
        if sl and sl.bullets:
            for item in sl.bullets[:4]:
                p = tf.add_paragraph()
                _runs(p, "•  " + item, size=size)
        _notes(slide, sl.notes if sl else None)

    def closing_slide(self) -> None:
        slide = _blank_slide(self.prs)
        self.number += 1
        _rect(slide, 0, 0, SLIDE_W, SLIDE_H, BACKGROUND)
        _rect(slide, 0, 0, Inches(0.35), SLIDE_H, PRIMARY)
        _textbox(slide, Inches(1.0), Inches(2.4), SLIDE_W - Inches(2.0), Inches(1.0), self.spec.closing_text or "Vielen Dank", size=34, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _textbox(slide, Inches(1.0), Inches(3.5), SLIDE_W - Inches(2.0), Inches(0.5), f"{self.spec.author} · {self.date}", size=14, color=MUTED)
        _textbox(slide, Inches(1.0), Inches(4.0), SLIDE_W - Inches(2.0), Inches(0.4), "Erstellt mit Vogelring-Analyse · vogelring.com", size=11, color=MUTED)

    def build(self) -> Presentation:
        self.title_slide()
        self.dataset_ids += [sl.dataset_id for sl in self.spec.slides if sl.dataset_id]
        for sl in self.spec.slides:
            for vid in sl.visual_ids:
                if self.visuals.exists(vid):
                    self.dataset_ids += self.visuals.meta(vid).dataset_ids
        explicit_methods = any(sl.kind == "methods" for sl in self.spec.slides)
        handlers = {
            "title": lambda sl: self.text_slide(Slide(kind="text", title=sl.title, subtitle=sl.subtitle, text=sl.text or "", notes=sl.notes)),
            "section": self.section_slide,
            "bullets": self.bullets_slide,
            "figure": self.figure_slide,
            "two_figures": self.two_figures_slide,
            "table": self.table_slide,
            "key_findings": self.key_findings_slide,
            "text": self.text_slide,
            "methods": self.methods_slide,
        }
        for sl in self.spec.slides:
            handlers[sl.kind](sl)
        if self.spec.include_methods and not explicit_methods and self.dataset_ids:
            self.methods_slide()
        self.closing_slide()
        return self.prs


def build_deck(spec: DeckSpec, *, visuals: VisualStore, datasets: DatasetStore, store: DeckStore) -> DeckMeta:
    """Write the PPTX for `spec`, register it in `store` and publish a deck.created event."""
    builder = _Builder(spec, visuals, datasets)
    try:
        prs = builder.build()
    except KeyError as exc:
        raise DeckBuildError(str(exc)) from exc
    visual_ids = list(dict.fromkeys(builder.visual_ids))
    dataset_ids = list(dict.fromkeys(builder.dataset_ids))
    meta = store.create(spec.model_dump(mode="json", exclude_none=True), {"title": spec.title, "slides": len(prs.slides), "visual_ids": visual_ids, "dataset_ids": dataset_ids})
    prs.save(str(store.dir(meta.id) / "deck.pptx"))
    meta.files = deck_urls(meta.id)
    store.save_meta(meta)
    bus.publish("deck.created", f"Präsentation erstellt: {spec.title}", deck_id=meta.id, deck_title=spec.title, slides=meta.slides, **meta.files)
    return meta


# ---------------------------------------------------------------- from report


def deck_from_report(report_model: dict[str, Any], *, max_bullets: int = 6) -> DeckSpec:
    """Derive a presentation from a built report model (reports.build_report_model output)."""
    slides: list[Slide] = []
    findings = [str(k) for k in report_model.get("key_findings") or []]
    if findings:
        slides.append(Slide(kind="key_findings", title="Das Wichtigste in Kürze", bullets=findings[:max_bullets]))
    for sec in report_model.get("sections") or []:
        heading = sec.get("heading") or "Ergebnisse"
        text = (sec.get("text") or "").strip()
        blocks = sec.get("blocks") or sec.get("items") or []
        figures = [b for b in blocks if b.get("type") == "figure" and b.get("id")]
        tables = [b for b in blocks if b.get("type") in ("table", "table_visual")]
        slides.append(Slide(kind="section", title=heading, subtitle=_first_sentence(text)))
        if text:
            bullets = _bullets_from_markdown(text, max_bullets)
            if bullets:
                slides.append(Slide(kind="bullets", title=heading, bullets=bullets))
        for fig in figures:
            slides.append(Slide(kind="figure", title=fig.get("caption") or heading, visual_ids=[fig["id"]], caption=f"Abbildung {fig['no']}" if fig.get("no") else None))
        for tab in tables:
            if tab.get("type") == "table_visual" and tab.get("id"):
                slides.append(Slide(kind="figure", title=tab.get("caption") or heading, visual_ids=[tab["id"]]))
            elif tab.get("dataset_id"):
                slides.append(Slide(kind="table", title=tab.get("caption") or heading, dataset_id=tab["dataset_id"]))
    recs = [str(r) for r in report_model.get("recommendations") or []]
    if recs:
        slides.append(Slide(kind="bullets", title="Empfehlungen", bullets=recs[:7]))
    if not slides:
        slides.append(Slide(kind="text", title=report_model.get("title") or "Bericht", text="Keine Inhalte im Bericht."))
    return DeckSpec(
        title=report_model.get("title") or "Bericht",
        subtitle=report_model.get("subtitle"),
        author=report_model.get("author") or "Vogelring-Analyse",
        date=report_model.get("date"),
        question=report_model.get("question"),
        slides=slides,
        include_methods=bool(report_model.get("include_methods", True)),
    )


def _first_sentence(text: str) -> str | None:
    plain = re.sub(r"[*_`#>|]", "", text or "").strip()
    if not plain:
        return None
    first = re.split(r"(?<=[.!?])\s", plain, maxsplit=1)[0]
    return first[:160] if first else None


def _bullets_from_markdown(text: str, limit: int) -> list[str]:
    bullets = [re.sub(r"^\s*(?:[-*•]|\d+[.)])\s+", "", ln).strip() for ln in text.splitlines() if re.match(r"^\s*(?:[-*•]|\d+[.)])\s+", ln)]
    if bullets:
        return bullets[:limit]
    plain = re.sub(r"\|.*\|", "", text)
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", re.sub(r"\s+", " ", plain)) if len(s.strip()) > 20]
    return sentences[:limit]
