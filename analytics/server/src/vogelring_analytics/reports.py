"""Report composer: markdown sections + visuals + tables → HTML, PDF (A4) and DOCX, with a plain-language
"So wurde gezählt" section and a technical appendix."""

from __future__ import annotations

import json
import re
import threading
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import markdown as md
import pandas as pd
from pydantic import BaseModel, ConfigDict, Field

from .datasets import DatasetStore, new_id
from .settings import settings
from .specs import ColumnSpec, TableSpec
from .tables import build_table_model
from .visuals import VisualStore


class ReportBuildError(ValueError):
    pass


class _Spec(BaseModel):
    model_config = ConfigDict(extra="forbid")


class ReportTable(_Spec):
    dataset_id: str
    caption: str | None = Field(None, description="Table caption (German).")
    columns: list[ColumnSpec] | None = None
    sort_by: str | None = None
    sort_desc: bool = True
    limit: int = Field(25, ge=1, le=200)
    totals_row: bool = False


class ReportSection(_Spec):
    heading: str = Field(..., description="Section title, e.g. 'Ergebnisse', 'Wo wurden die Vögel gesehen?'")
    text: str = Field("", description="Markdown body in plain German. Explain what the reader sees. **bold**, lists and simple tables are fine.")
    visual_ids: list[str] = Field(default_factory=list, description="Charts/maps/tables (visual_id) to place after the text, each with a numbered caption.")
    captions: dict[str, str] | None = Field(None, description="visual_id → caption override (default: the visual's title).")
    tables: list[ReportTable] = Field(default_factory=list, description="Datasets rendered as formatted tables.")


class ReportSpec(_Spec):
    title: str
    subtitle: str | None = None
    author: str = Field("Vogelring-Analyse", description="Shown under the title, e.g. 'Ingo R. mit Vogelring-Analyse'.")
    question: str | None = Field(None, description="The question the report answers, in one sentence (shown at the top).")
    key_findings: list[str] = Field(default_factory=list, description="3–6 plain sentences with the main results; printed as a box at the top.")
    sections: list[ReportSection] = Field(..., min_length=1)
    include_methods: bool = Field(True, description="Append 'So wurde gezählt': plain-language data basis per dataset (filters, period, counts).")
    include_sql_appendix: bool = Field(True, description="Append the technical appendix (SQL) for readers who want to check.")
    methods_note: str | None = Field(None, description="Extra plain-language notes on assumptions (e.g. 'Mausergast = Mai–Juli im Ostpark').")
    recommendations: list[str] = Field(default_factory=list, description="Optional: what to do next (observe more at X, check records Y).")


@dataclass
class ReportMeta:
    id: str
    title: str
    created_at: str
    visual_ids: list[str]
    dataset_ids: list[str]
    sections: int
    files: dict[str, str] = field(default_factory=dict)


_MD = md.Markdown(extensions=["tables", "sane_lists", "smarty"], output_format="html5")


def md_html(text: str) -> str:
    _MD.reset()
    return _MD.convert(text or "")


class ReportStore:
    def __init__(self, root: Path | None = None) -> None:
        self.root = root or settings.reports_dir
        self.root.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()

    def dir(self, rid: str) -> Path:
        if not rid or any(ch in rid for ch in "/\\.."):
            raise KeyError(f"Invalid report id: {rid!r}")
        return self.root / rid

    def create(self, spec: dict[str, Any], model: dict[str, Any], meta_kwargs: dict[str, Any]) -> ReportMeta:
        rid = new_id("rp_")
        d = self.dir(rid)
        d.mkdir(parents=True, exist_ok=False)
        meta = ReportMeta(id=rid, created_at=datetime.now(timezone.utc).isoformat(timespec="seconds"), **meta_kwargs)
        with self._lock:
            (d / "spec.json").write_text(json.dumps(spec, ensure_ascii=False, indent=1, default=str))
            (d / "model.json").write_text(json.dumps(model, ensure_ascii=False, default=str))
            (d / "meta.json").write_text(json.dumps(asdict(meta), ensure_ascii=False, indent=1))
        return meta

    def meta(self, rid: str) -> ReportMeta:
        p = self.dir(rid) / "meta.json"
        if not p.exists():
            raise KeyError(f"Unknown report id: {rid}")
        return ReportMeta(**json.loads(p.read_text()))

    def model(self, rid: str) -> dict[str, Any]:
        return json.loads((self.dir(rid) / "model.json").read_text())

    def exists(self, rid: str) -> bool:
        try:
            return (self.dir(rid) / "meta.json").exists()
        except KeyError:
            return False

    def write_bytes(self, rid: str, name: str, data: bytes) -> Path:
        p = self.dir(rid) / name
        p.write_bytes(data)
        return p

    def file(self, rid: str, name: str) -> Path | None:
        p = self.dir(rid) / name
        return p if p.exists() else None

    def list(self, limit: int = 20) -> list[ReportMeta]:
        metas = []
        for p in self.root.glob("rp_*/meta.json"):
            try:
                metas.append(ReportMeta(**json.loads(p.read_text())))
            except Exception:
                continue
        metas.sort(key=lambda m: m.created_at, reverse=True)
        return metas[:limit]


# ---------------------------------------------------------------- model building


def _period(df: pd.DataFrame) -> str | None:
    for c in ("date", "datum", "ringing_date", "beringung"):
        if c in df.columns and pd.api.types.is_datetime64_any_dtype(df[c]) and df[c].notna().any():
            return f"{df[c].min().strftime('%d.%m.%Y')} bis {df[c].max().strftime('%d.%m.%Y')}"
    for c in ("jahr", "year", "beringungsjahr"):
        if c in df.columns and df[c].notna().any():
            try:
                return f"{int(df[c].min())}–{int(df[c].max())}"
            except Exception:
                pass
    return None


def _dataset_facts(ds_id: str, datasets: DatasetStore) -> dict[str, Any]:
    meta = datasets.meta(ds_id)
    df = datasets.load(ds_id)
    facts: dict[str, Any] = {"id": ds_id, "name": meta.name, "description": meta.description, "rows": meta.row_count, "sql": meta.sql, "created_at": meta.created_at[:10], "source": meta.source}
    facts["period"] = _period(df)
    if "ring" in df.columns:
        facts["birds"] = int(df["ring"].nunique())
    if "place" in df.columns:
        facts["places"] = int(df["place"].nunique())
    if "species" in df.columns and df["species"].notna().any():
        vc = df["species"].value_counts()
        facts["species"] = ", ".join(f"{k} ({v})" for k, v in vc.head(5).items()) + (" …" if len(vc) > 5 else "")
    return facts


def build_report_model(spec: ReportSpec, visuals: VisualStore, datasets: DatasetStore) -> tuple[dict[str, Any], dict[str, Any]]:
    fig_no = 0
    tab_no = 0
    sections = []
    dataset_ids: list[str] = []
    visual_ids: list[str] = []
    for sec in spec.sections:
        blocks: list[dict[str, Any]] = []
        for vid in sec.visual_ids:
            if not visuals.exists(vid):
                raise ReportBuildError(f"Unbekannte visual_id {vid}.")
            vm = visuals.meta(vid)
            model = visuals.model(vid)
            visual_ids.append(vid)
            dataset_ids += vm.dataset_ids
            caption = (sec.captions or {}).get(vid) or vm.title or vm.kind
            if vm.kind == "table":
                tab_no += 1
                blocks.append({"type": "table_visual", "id": vid, "no": tab_no, "caption": caption, "model": model})
            else:
                fig_no += 1
                has_png = visuals.file(vid, "image.png") is not None
                blocks.append({"type": "figure", "id": vid, "no": fig_no, "caption": caption, "kind": vm.kind, "width": vm.width, "height": vm.height, "png": f"/viz/{vid}.png" if has_png else None, "interactive": f"{settings.public_base_url}/viz/{vid}"})
        for t in sec.tables:
            if not datasets.exists(t.dataset_id):
                raise ReportBuildError(f"Unbekannte dataset_id {t.dataset_id}.")
            dataset_ids.append(t.dataset_id)
            tspec = TableSpec(dataset_id=t.dataset_id, columns=t.columns, sort_by=t.sort_by, sort_desc=t.sort_desc, limit=t.limit, totals_row=t.totals_row, row_number=False)
            tm, _ = build_table_model(tspec, datasets.load(t.dataset_id))
            tab_no += 1
            blocks.append({"type": "table", "no": tab_no, "caption": t.caption or datasets.meta(t.dataset_id).name, "model": tm})
        sections.append({"heading": sec.heading, "html": md_html(sec.text), "text": sec.text, "blocks": blocks})

    seen: list[str] = []
    for d in dataset_ids:
        if d not in seen and datasets.exists(d):
            seen.append(d)
    facts = [_dataset_facts(d, datasets) for d in seen]
    now = datetime.now()
    model = {
        "title": spec.title,
        "subtitle": spec.subtitle,
        "author": spec.author,
        "question": spec.question,
        "date": now.strftime("%d.%m.%Y"),
        "key_findings": spec.key_findings,
        "sections": sections,
        "recommendations": spec.recommendations,
        "include_methods": spec.include_methods,
        "include_sql_appendix": spec.include_sql_appendix,
        "methods_note_html": md_html(spec.methods_note) if spec.methods_note else None,
        "methods_note": spec.methods_note,
        "datasets": facts,
        "figures": fig_no,
        "tables": tab_no,
        "base_url": settings.public_base_url,
    }
    summary = {"sections": len(sections), "figures": fig_no, "tables": tab_no, "datasets": len(facts)}
    return model, summary


# ---------------------------------------------------------------- DOCX


def _add_markdown(doc, text: str) -> None:
    from docx.shared import Pt

    lines = (text or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-+:?$", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Light Grid Accent 1"
                for r_i, r in enumerate(rows):
                    for c_i in range(ncol):
                        cell = t.cell(r_i, c_i)
                        cell.text = r[c_i] if c_i < len(r) else ""
                        if r_i == 0:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.bold = True
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(4, len(m.group(1)) + 1))
            i += 1
            continue
        m = re.match(r"^\s*[-*•]\s+(.*)", line)
        if m:
            _runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if m:
            _runs(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            _runs(p, line.lstrip("> "))
            for r in p.runs:
                r.italic = True
            i += 1
            continue
        # paragraph: merge following lines until blank
        buf = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s|\s*[-*•]\s|\s*\d+[.)]\s|\||>)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        _runs(p, " ".join(buf))
        p.paragraph_format.space_after = Pt(6)


def _runs(p, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            p.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            p.add_run(part[1:-1]).font.name = "Consolas"
        else:
            p.add_run(part)


def _docx_table(doc, tm: dict[str, Any], caption: str, no: int) -> None:
    from docx.shared import Pt

    cols = tm["columns"]
    rows = tm["rows"][:60]
    t = doc.add_table(rows=len(rows) + 1 + (1 if tm.get("totals") else 0), cols=len(cols))
    t.style = "Light Grid Accent 1"
    for j, c in enumerate(cols):
        cell = t.cell(0, j)
        cell.text = c["label"]
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
    for i, row in enumerate(rows, start=1):
        for j, cell in enumerate(row):
            tc = t.cell(i, j)
            tc.text = cell["text"]
            for p in tc.paragraphs:
                if cols[j]["align"] == "right":
                    p.alignment = 2
                for r in p.runs:
                    r.font.size = Pt(9)
    if tm.get("totals"):
        for j, v in enumerate(tm["totals"]):
            tc = t.cell(len(rows) + 1, j)
            tc.text = v
            for p in tc.paragraphs:
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9)
    cap = doc.add_paragraph()
    r = cap.add_run(f"Tabelle {no}: {caption}")
    r.italic = True
    r.font.size = Pt(9)
    if tm.get("total_rows", 0) > len(rows):
        doc.add_paragraph(f"Gezeigt: {len(rows)} von {tm['total_rows']} Zeilen.").runs[0].font.size = Pt(8)


def write_docx(model: dict[str, Any], path: Path, visuals: VisualStore) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    h = doc.add_heading(model["title"], level=0)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x00, 0x43, 0x6C)
    if model.get("subtitle"):
        p = doc.add_paragraph()
        r = p.add_run(model["subtitle"])
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x5F, 0x6B, 0x73)
    meta = doc.add_paragraph()
    r = meta.add_run(f"{model['author']} · {model['date']}")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x5F, 0x6B, 0x73)
    if model.get("question"):
        p = doc.add_paragraph()
        p.add_run("Fragestellung: ").bold = True
        p.add_run(model["question"])
    if model.get("key_findings"):
        doc.add_heading("Das Wichtigste in Kürze", level=1)
        for k in model["key_findings"]:
            _runs(doc.add_paragraph(style="List Bullet"), k)

    for sec in model["sections"]:
        doc.add_heading(sec["heading"], level=1)
        _add_markdown(doc, sec["text"])
        for it in sec["blocks"]:
            if it["type"] == "figure" and it.get("png"):
                png = visuals.file(it["id"], "image.png")
                if png:
                    doc.add_picture(str(png), width=Cm(16))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = cap.add_run(f"Abbildung {it['no']}: {it['caption']}")
                rr.italic = True
                rr.font.size = Pt(9)
                link = doc.add_paragraph()
                link.alignment = WD_ALIGN_PARAGRAPH.CENTER
                lr = link.add_run(f"Interaktiv: {it['interactive']}")
                lr.font.size = Pt(8)
                lr.font.color.rgb = RGBColor(0x5F, 0x6B, 0x73)
            elif it["type"] in ("table", "table_visual"):
                _docx_table(doc, it["model"], it["caption"], it["no"])

    if model.get("recommendations"):
        doc.add_heading("Empfehlungen", level=1)
        for k in model["recommendations"]:
            _runs(doc.add_paragraph(style="List Bullet"), k)

    if model.get("include_methods"):
        doc.add_heading("So wurde gezählt", level=1)
        doc.add_paragraph(
            "Grundlage sind die Sichtungen und Beringungen in der Vogelring-Datenbank (Stand " + model["date"] + "). "
            "Jede Sichtung ist eine Ablesung eines beringten Vogels an einem Ort an einem Tag. Nicht gesehen heißt nicht abwesend: "
            "die Daten zeigen vor allem, wo und wann beobachtet wurde."
        )
        if model.get("methods_note"):
            _add_markdown(doc, model["methods_note"])
        for i, d in enumerate(model["datasets"], start=1):
            p = doc.add_paragraph(style="List Number")
            p.add_run(d["name"]).bold = True
            bits = []
            if d.get("description"):
                bits.append(d["description"])
            bits.append(f"{d['rows']} Zeilen")
            if d.get("period"):
                bits.append(f"Zeitraum {d['period']}")
            if d.get("birds") is not None:
                bits.append(f"{d['birds']} verschiedene Vögel")
            if d.get("places") is not None:
                bits.append(f"{d['places']} Orte")
            if d.get("species"):
                bits.append(f"Arten: {d['species']}")
            p.add_run(" – " + "; ".join(bits))

    if model.get("include_sql_appendix"):
        doc.add_heading("Anhang: technische Details", level=1)
        doc.add_paragraph("Für alle, die nachrechnen möchten: die Datenbankabfragen hinter den Zahlen.")
        for d in model["datasets"]:
            doc.add_paragraph(d["name"]).runs[0].bold = True
            p = doc.add_paragraph()
            r = p.add_run(d["sql"].strip())
            r.font.name = "Consolas"
            r.font.size = Pt(8)
    doc.save(path)
    return path
