"""Data exports: formatted XLSX and CSV."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

_HEADER_FILL = PatternFill("solid", fgColor="00436C")
_HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_BORDER = Border(bottom=Side(style="thin", color="DDDDDD"))


def safe_filename(name: str, default: str = "export") -> str:
    name = re.sub(r"[^\w\-. ]+", "_", (name or "").strip(), flags=re.UNICODE).strip(" ._")
    name = re.sub(r"[\s_]+", "_", name)
    return name[:80] or default


def write_csv(df: pd.DataFrame, path: Path) -> Path:
    df.to_csv(path, index=False, sep=";", date_format="%Y-%m-%d", encoding="utf-8-sig")
    return path


def write_xlsx(df: pd.DataFrame, path: Path, *, sheet_name: str = "Daten", title: str | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] or "Daten"
    start_row = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14, color="00436C")
        start_row = 3

    cols = [str(c) for c in df.columns]
    for j, name in enumerate(cols, start=1):
        cell = ws.cell(row=start_row, column=j, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    date_cols = {c for c in df.columns if pd.api.types.is_datetime64_any_dtype(df[c])}
    for i, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for j, c in enumerate(df.columns, start=1):
            v = row[c]
            if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
                continue
            if c in date_cols:
                cell = ws.cell(row=i, column=j, value=pd.Timestamp(v).to_pydatetime())
                cell.number_format = "DD.MM.YYYY"
            elif hasattr(v, "item"):
                cell = ws.cell(row=i, column=j, value=v.item())
            else:
                cell = ws.cell(row=i, column=j, value=v)
            cell.border = _BORDER

    # Column widths from content (capped)
    for j, c in enumerate(df.columns, start=1):
        sample = df[c].astype(str).head(500).map(len)
        width = min(max([len(str(c))] + sample.tolist(), default=10) + 2, 50)
        ws.column_dimensions[get_column_letter(j)].width = max(width, 8)

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if len(df) > 0 and len(cols) > 0:
        ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + len(df)}"
        table = Table(displayName=re.sub(r"\W", "", f"T{ws.title}") or "Tabelle", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(table)
    wb.save(path)
    return path
